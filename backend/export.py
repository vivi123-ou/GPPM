"""Xuất tài liệu kiểm thử: test case (Excel/CSV) và báo cáo kết quả chạy (Excel).

File Excel định dạng sẵn để đưa thẳng vào tài liệu dự án:
- Tổng quan: thống kê theo loại / độ ưu tiên, model AI, thời điểm xuất
- Test cases: mỗi bước một dòng (bước, hành động, đối tượng, dữ liệu, kết quả mong đợi)
- Dữ liệu test: từng data set (phân vùng hợp lệ / không hợp lệ)
- Yêu cầu: các yêu cầu gốc đã nhập
- Kết quả (báo cáo): pass/fail, lỗi, nhãn AI và giải thích
"""
import csv
import io
import json
from collections import Counter
from datetime import datetime

from fastapi import APIRouter
from fastapi.responses import Response
from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

from backend.db import get_conn
from backend.models import Step, TestCase

router = APIRouter(prefix="/api/export", tags=["Export"])

XLSX = "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
VIOLET, VIOLET_LIGHT = "6D28D9", "F3EEFF"
GREEN_LIGHT, RED_LIGHT, AMBER_LIGHT = "DCFCE7", "FEE2E2", "FEF3C7"
THIN = Side(style="thin", color="DDD6FE")
VERDICT_VI = {"app_bug": "Lỗi ứng dụng", "test_broken": "Lỗi kịch bản", "flaky": "Không ổn định"}
PRIORITY_VI = {"high": "Cao", "medium": "Trung bình", "low": "Thấp"}


def _load(ids: str | None):
    conn = get_conn()
    rows = conn.execute("SELECT id, requirement_id, json, version FROM test_cases "
                        "WHERE status='active' ORDER BY id").fetchall()
    reqs = conn.execute("SELECT id, text, created_at FROM requirements ORDER BY id").fetchall()
    conn.close()
    wanted = {i.strip() for i in ids.split(",")} if ids else None
    cases = [(TestCase.model_validate_json(r["json"]), r["version"]) for r in rows
             if not wanted or r["id"] in wanted]
    return cases, reqs


def _j(v) -> str:
    return "" if v in (None, "", {}, []) else json.dumps(v, ensure_ascii=False)


def describe_step(s: Step) -> tuple[str, str, str, str]:
    """(hành động, đối tượng, dữ liệu, kết quả mong đợi) dễ đọc cho người."""
    if s.method:
        exp = s.expect if isinstance(s.expect, dict) else {}
        expected = []
        if exp.get("status") is not None:
            expected.append(f"HTTP {exp['status']}")
        if exp.get("json_contains"):
            expected.append("response chứa " + _j(exp["json_contains"]))
        if s.save:
            expected.append("lưu " + ", ".join(f"{k} ← {v}" for k, v in s.save.items()))
        return s.method, s.path or "", _j(s.body), "; ".join(expected)
    t = s.target
    target = ""
    if t:
        target = (f'testid="{t.testid}"' if t.testid else
                  f'{t.role or ""} "{t.name}"'.strip() if t.name else t.css or "")
    if s.action == "goto":
        target = s.url or ""
    expected = f'hiển thị "{s.expect}"' if s.action == "assert_text" else \
        "phần tử hiển thị" if s.action == "assert_visible" else ""
    return s.action or "", target, s.value or "", expected


def _header(ws, row: int, headers: list[str]):
    for c, h in enumerate(headers, 1):
        cell = ws.cell(row=row, column=c, value=h)
        cell.font = Font(bold=True, color="FFFFFF")
        cell.fill = PatternFill("solid", fgColor=VIOLET)
        cell.alignment = Alignment(vertical="center", wrap_text=True)
        cell.border = Border(top=THIN, bottom=THIN, left=THIN, right=THIN)
    ws.freeze_panes = ws.cell(row=row + 1, column=1)
    ws.auto_filter.ref = f"A{row}:{get_column_letter(len(headers))}{row}"


def _rows(ws, start: int, rows: list[list], widths: list[int], fills: list[str | None] | None = None):
    for i, values in enumerate(rows):
        for c, v in enumerate(values, 1):
            cell = ws.cell(row=start + i, column=c, value=v)
            cell.alignment = Alignment(vertical="top", wrap_text=True)
            cell.border = Border(top=THIN, bottom=THIN, left=THIN, right=THIN)
            if fills and fills[i]:
                cell.fill = PatternFill("solid", fgColor=fills[i])
    for c, w in enumerate(widths, 1):
        ws.column_dimensions[get_column_letter(c)].width = w


def _title(ws, title: str, subtitle: str):
    ws["A1"] = title
    ws["A1"].font = Font(bold=True, size=16, color=VIOLET)
    ws["A2"] = subtitle
    ws["A2"].font = Font(italic=True, color="6B7280")


def _model_name() -> str:
    try:
        from backend import ai_client
        if ai_client.USE_MOCK:
            return "Mock (USE_MOCK=1)"
        from ai.llm_provider import get_provider
        p = get_provider()
        return f"{p.name} / {p.model}"
    except Exception:
        return "-"


def _send(wb: Workbook, name: str) -> Response:
    buf = io.BytesIO()
    wb.save(buf)
    fname = f"{name}_{datetime.now():%Y%m%d_%H%M}.xlsx"
    return Response(buf.getvalue(), media_type=XLSX,
                    headers={"Content-Disposition": f'attachment; filename="{fname}"'})


@router.get("/test-cases.xlsx")
def export_test_cases_xlsx(ids: str | None = None):
    cases, reqs = _load(ids)
    wb = Workbook()
    now = f"{datetime.now():%d/%m/%Y %H:%M}"

    ws = wb.active
    ws.title = "Tổng quan"
    _title(ws, "TestPilot – Tài liệu kịch bản kiểm thử", f"Xuất lúc {now} · Mô hình AI: {_model_name()}")
    types = Counter(tc.type for tc, _ in cases)
    prio = Counter(tc.priority for tc, _ in cases)
    summary = [["Tổng số test case", len(cases)],
               ["Test API", types.get("api", 0)], ["Test UI", types.get("ui", 0)],
               ["Tổng số bước", sum(len(tc.steps) for tc, _ in cases)],
               ["Tổng số bộ dữ liệu", sum(len(tc.data_sets) for tc, _ in cases)],
               ["Ưu tiên cao", prio.get("high", 0)], ["Ưu tiên trung bình", prio.get("medium", 0)],
               ["Ưu tiên thấp", prio.get("low", 0)], ["Số yêu cầu", len(reqs)]]
    _header(ws, 4, ["Chỉ số", "Giá trị"])
    _rows(ws, 5, summary, [32, 18])
    ws.freeze_panes = None

    ws = wb.create_sheet("Test cases")
    _header(ws, 1, ["ID", "Tiêu đề", "Loại", "Ưu tiên", "Yêu cầu", "Phiên bản", "Bước",
                    "Hành động / Method", "Đối tượng / Endpoint", "Dữ liệu / Body", "Kết quả mong đợi"])
    rows, fills = [], []
    for n, (tc, version) in enumerate(cases):
        fill = VIOLET_LIGHT if n % 2 == 0 else None
        for i, s in enumerate(tc.steps, 1):
            first = i == 1
            rows.append([tc.id if first else "", tc.title if first else "",
                         tc.type.upper() if first else "", PRIORITY_VI.get(tc.priority, tc.priority) if first else "",
                         tc.requirement_id or "" if first else "", version if first else "",
                         i, *describe_step(s)])
            fills.append(fill)
    _rows(ws, 2, rows, [14, 38, 7, 11, 10, 9, 6, 16, 34, 40, 40], fills)

    ws = wb.create_sheet("Dữ liệu test")
    _header(ws, 1, ["ID test case", "Tên bộ dữ liệu", "Kỳ vọng", "Giá trị"])
    rows, fills = [], []
    for tc, _ in cases:
        for d in tc.data_sets:
            rows.append([tc.id, d.name, "Hợp lệ (accept)" if d.expected == "accept" else "Không hợp lệ (reject)",
                         _j(d.values)])
            fills.append(GREEN_LIGHT if d.expected == "accept" else RED_LIGHT)
    _rows(ws, 2, rows, [16, 28, 20, 70], fills)

    ws = wb.create_sheet("Yêu cầu")
    _header(ws, 1, ["Mã", "Thời điểm", "Nội dung yêu cầu"])
    _rows(ws, 2, [[f"REQ-{r['id']}", r["created_at"], r["text"]] for r in reqs], [10, 20, 100])
    return _send(wb, "TestPilot_TestCases")


@router.get("/test-cases.csv")
def export_test_cases_csv(ids: str | None = None):
    cases, _ = _load(ids)
    buf = io.StringIO()
    w = csv.writer(buf)
    w.writerow(["ID", "Tiêu đề", "Loại", "Ưu tiên", "Bước", "Hành động", "Đối tượng", "Dữ liệu", "Kết quả mong đợi"])
    for tc, _ in cases:
        for i, s in enumerate(tc.steps, 1):
            w.writerow([tc.id, tc.title, tc.type, tc.priority, i, *describe_step(s)])
    # BOM để Excel mở đúng tiếng Việt
    return Response("﻿" + buf.getvalue(), media_type="text/csv; charset=utf-8",
                    headers={"Content-Disposition": f'attachment; filename="TestPilot_TestCases_{datetime.now():%Y%m%d_%H%M}.csv"'})


@router.post("/report.xlsx")
def export_report_xlsx(report: dict):
    """Nhận đúng JSON trả về từ POST /api/ai/run và xuất thành báo cáo Excel."""
    results = report.get("results") or []
    wb = Workbook()
    ws = wb.active
    ws.title = "Tổng quan"
    total, passed = len(results), sum(r.get("status") == "passed" for r in results)
    verdicts = Counter((r.get("ai_analysis") or {}).get("verdict") for r in results if r.get("status") == "failed")
    _title(ws, "TestPilot – Báo cáo kết quả kiểm thử",
           f"Xuất lúc {datetime.now():%d/%m/%Y %H:%M} · Hệ thống: {report.get('target_url', '-')} · AI: {_model_name()}")
    _header(ws, 4, ["Chỉ số", "Giá trị"])
    _rows(ws, 5, [["Tổng lượt chạy", total], ["Đạt", passed], ["Không đạt", total - passed],
                  ["Tỉ lệ đạt", f"{(passed / total * 100 if total else 0):.0f}%"],
                  ["Lỗi ứng dụng (app_bug)", verdicts.get("app_bug", 0)],
                  ["Lỗi kịch bản (test_broken)", verdicts.get("test_broken", 0)],
                  ["Không ổn định (flaky)", verdicts.get("flaky", 0)]], [32, 40])
    ws.freeze_panes = None

    ws = wb.create_sheet("Kết quả")
    _header(ws, 1, ["ID test case", "Tiêu đề", "Bộ dữ liệu", "Kết quả", "Thời gian (ms)", "Request",
                    "HTTP mong đợi", "HTTP thực tế", "Lỗi", "AI phân loại", "Độ tin cậy", "Giải thích của AI",
                    "Đề xuất sửa"])
    rows, fills = [], []
    for r in results:
        a = r.get("ai_analysis") or {}
        req = r.get("request") or {}
        ok = r.get("status") == "passed"
        rows.append([r.get("test_case_id"), r.get("title"), r.get("data_set"), "ĐẠT" if ok else "KHÔNG ĐẠT",
                     r.get("duration_ms"), f"{req.get('method', '')} {req.get('path', '')}".strip(),
                     r.get("expected_status"), r.get("actual_status"), r.get("error_msg") or "",
                     VERDICT_VI.get(a.get("verdict"), a.get("verdict") or ""),
                     f"{a['confidence']:.0%}" if isinstance(a.get("confidence"), (int, float)) else "",
                     a.get("explanation") or "", a.get("suggested_fix") or ""])
        fills.append(GREEN_LIGHT if ok else (AMBER_LIGHT if a.get("verdict") == "test_broken" else RED_LIGHT))
    _rows(ws, 2, rows, [14, 36, 18, 8, 10, 30, 10, 10, 36, 16, 10, 60, 40], fills)
    return _send(wb, "TestPilot_BaoCao")
